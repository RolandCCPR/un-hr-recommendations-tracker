"""Write the implementation assessment back into a country's OHCHR thematic-list
.docx, filling the empty 4th column ("Assessment/comments on level of
implementation") with the grade, the reasoned assessment and the quotes.

    python -m scripts.fill_thematic_list [slug]     (default: united-states)

Requires python-docx (a helper-script dependency, not used by the package).
"""

from __future__ import annotations

import re
import shutil
import sys
import tempfile
from pathlib import Path

import docx
from docx.opc.exceptions import PackageNotFoundError

from recommendations.repository import RecommendationRepository

from scripts.countries import ROOT, get, out_dir

_PARA_RE = re.compile(r"^\s*(\d+\.\d+)\b")


def _open(path: Path):
    try:
        return docx.Document(str(path))
    except (PermissionError, PackageNotFoundError):
        if not path.exists():
            raise
        tmp = Path(tempfile.gettempdir()) / "thematic_src_copy.docx"
        shutil.copy2(path, tmp)
        return docx.Document(str(tmp))


def build(slug: str) -> None:
    cfg = get(slug)
    repo = RecommendationRepository(cfg["db"])
    by_para = {r.paragraph: r for r in repo.list()}

    document = _open(cfg["thematic"])
    filled = 0
    for table in document.tables:
        for row in table.rows:
            cells = row.cells
            if len(cells) < 4:
                continue
            m = _PARA_RE.match(cells[0].text.strip())
            if not m:
                continue
            rec = by_para.get(m.group(1))
            if rec is None:
                continue

            target = cells[3]
            target.text = ""
            p = target.paragraphs[0]
            if rec.action_summary:
                p.add_run(f"Action recommended: {rec.action_summary}\n").italic = True
            p.add_run(f"Grade {rec.grade.value} — {rec.grade.committee_meaning}").bold = True
            p.add_run(f"\nPosition at review: {rec.position.value.replace('_', '-')}.")
            p.add_run("\n\nAssessment: ").bold = True
            p.add_run(rec.assessment_rationale or "")
            if rec.assessment_evidence:
                p.add_run("\n\nEvidence:").bold = True
                for line in rec.assessment_evidence.splitlines():
                    if line.strip():
                        p.add_run(f"\n• {line.strip()}")
            filled += 1

    dst = out_dir(slug) / cfg["thematic_out"]
    document.save(str(dst))
    print(f"Filled {filled} rows -> {dst.relative_to(ROOT)}")
    repo.close()


if __name__ == "__main__":
    for s in sys.argv[1:] or ["united-states"]:
        build(s)
